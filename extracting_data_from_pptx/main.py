import copy
import math
import os
import re
import sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData, ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from ppt_to_excel import pptx_tables_to_excel


EXCEL_NAME = "pptx_tables.xlsx"
REVIEW_NAME = "chart_data_review.xlsx"
OUTPUT_PPTX = "program_highlights_generated.pptx"
TEMPLATE_PPTX = "program_highlights.pptx"

# Excel sheet name (= pptx file name) -> pillar title shown on the slide
PILLARS = {
    "lag": "LEGAL AID AND GOVERNANCE",
    "mhpss": "MENTAL HEALTH and PSYCHOSOCIAL SUPPORT",
    "wge": "WOMEN & GIRLS",
    "psj": "PEACE AND SOCIAL JUSTICE",
}

# Optional: replace an auto-generated label with a short one.
#   key = start of the auto label (case-insensitive), value = label to use
LABEL_OVERRIDES = {
    # "# of DPHCs who receive legal aid": "# of DPHCs who receive legal aid and counselling",
}

MAX_LABEL_CHARS = 28             # axis labels of COLUMN charts are cut to this (bar charts wrap)
DECK_TITLE = "PROGRAMMES HIGHLIGHTS BASED ON ANNUAL BUSINESS PLANS"

# Colours / fonts (match the "Wood Type" theme of program_highlights.pptx)
C_TARGET = RGBColor(0xD3, 0x48, 0x17)
C_ACHIEVED = RGBColor(0x9B, 0x2D, 0x1F)
C_GREY = RGBColor(0x59, 0x59, 0x59)
FONT = "Garamond"
CHART_FONT_SIZE = 12


NUM = r"\d[\d,]*(?:\.\d+)?"
MALE_TOKENS = {"m", "male", "males", "men", "man", "b", "boy", "boys"}
FEMALE_TOKENS = {"f", "female", "females", "women", "woman", "w", "g", "girl", "girls"}


def to_num(text):
    return float(text.replace(",", "")) if text else None


def clean(value):
    return "" if value is None else str(value).replace("\xa0", " ").strip()


def split_groups(text):
    """A cell holding several items separated by blank lines -> list of items."""
    return [g.strip() for g in re.split(r"\n\s*\n+", text) if g.strip()]


def parse_target(text):
    """Return (label_in_target_cell, number)."""
    text = clean(text)
    if ":" in text and re.search(NUM, text.rsplit(":", 1)[1]):
        label, rest = text.rsplit(":", 1)       # 'Training ... Act, 2022: 2' -> 2
        m = re.search(NUM, rest)
        number = to_num(m.group())
    else:
        m = re.search(NUM, text)
        if not m:
            return "", None
        label, number = text[: m.start()], to_num(m.group())
    label = re.sub(r"^\s*target\s*:?", "", label, flags=re.I).strip(" :\n-")
    return label, number


def parse_achievement(text):
    """Return (achieved_number, male, female) from an 'Achievement' cell."""
    text = clean(text)
    if not text:
        return None, 0, 0

    segments = []                                        # text where the figures live
    total = re.search(r"Total[^\d\n]{0,15}?(" + NUM + ")", text, flags=re.I)
    if total:                                            # 'Total - 342(183M, 100F ...'
        number = to_num(total.group(1))
        segments = [text[total.start():].split("\n")[0]]
    else:
        quarters = re.findall(r"Q\d\s*[-–—:]\s*(" + NUM + ")", text)
        if quarters:                                     # sum of Q1 + Q2 ...
            number = sum(to_num(q) for q in quarters)
            segments = [ln for ln in text.split("\n") if re.match(r"\s*Q\d", ln)]
        else:
            lead = re.match(r"\s*(" + NUM + ")", text)   # '92 .' / '181 (111 female ...'
            if not lead:
                return None, 0, 0
            number = to_num(lead.group(1))
            segments = [text.split("\n")[0]]

    male = female = 0
    for seg in segments:
        for count, tag in re.findall(r"(" + NUM + r")\s*([A-Za-z]+)\b", seg):
            tag = tag.lower()
            if tag in MALE_TOKENS:
                male += to_num(count)
            elif tag in FEMALE_TOKENS:
                female += to_num(count)
    return number, male, female


def shorten(label):
    label = re.sub(r"\s+", " ", label).strip(" .:;-")
    label = re.sub(r"^#\s*\d[\d,]*\s+", "# ", label)      # '# 7500 of x' -> '# of x'
    for start, new in LABEL_OVERRIDES.items():
        if label.lower().startswith(start.lower()):
            return new
    if len(label) > MAX_LABEL_CHARS:
        cut = label[:MAX_LABEL_CHARS].rsplit(" ", 1)[0]
        label = cut.rstrip(" ,;:-") + "…"
    return label


def find_header(rows):
    """Locate column positions of indicator / target / achievement in a table."""
    for i, row in enumerate(rows):
        up = [clean(c).upper() for c in row]
        t = next((j for j, c in enumerate(up) if "TARGET" in c), None)
        a = next((j for j, c in enumerate(up)
                  if "ACHIEVEMENT" in c or "INCREAMENTAL" in c), None)
        if t is not None and a is not None:
            ind = next((j for j, c in enumerate(up)
                        if c.startswith(("INDICATOR", "ACTIVIT"))), 0)
            return i, ind, t, a
    return None


def sheet_to_tables(ws):
    """Split a sheet into the blocks written by pptx_tables_to_excel."""
    blocks, current = [], None
    for row in ws.iter_rows(values_only=True):
        first = clean(row[0]) if row else ""
        if re.match(r"Slide \d+ - Table \d+", first):
            current = {"label": first, "rows": []}
            blocks.append(current)
        elif current is not None and any(clean(c) for c in row):
            current["rows"].append(row)
    return blocks


def extract_indicators(excel_path, pillars=PILLARS):
    """Read every pillar sheet -> list of dict records."""
    wb = load_workbook(excel_path, data_only=True)
    records = []
    for sheet in wb.sheetnames:
        if sheet not in pillars:
            print(f"  (sheet '{sheet}' is not in PILLARS - skipped)")
            continue
        for block in sheet_to_tables(wb[sheet]):
            head = find_header(block["rows"])
            if not head:                       # challenge tables, narratives, etc.
                continue
            h_row, c_ind, c_tgt, c_ach = head
            for row in block["rows"][h_row + 1:]:
                cells = [clean(c) for c in row] + [""] * 6
                ind, tgt, ach = cells[c_ind], cells[c_tgt], cells[c_ach]
                if not tgt or not ach:
                    continue
                i_g, t_g, a_g = split_groups(ind), split_groups(tgt), split_groups(ach)
                if len(i_g) > 1 and len(i_g) == len(t_g) == len(a_g):
                    triples = zip(i_g, t_g, a_g)          # several indicators in 1 cell
                else:
                    triples = [(ind, tgt, ach)]
                for i_txt, t_txt, a_txt in triples:
                    t_label, target = parse_target(t_txt)
                    achieved, male, female = parse_achievement(a_txt)
                    if target is None or achieved is None:
                        continue
                    label = t_label if (t_label and not i_txt.strip()) else i_txt
                    if t_label and len(split_groups(ind)) != len(t_g):
                        label = t_label                   # label lives in target cell
                    records.append({
                        "pillar": sheet, "source": block["label"],
                        "indicator": shorten(label), "target": target,
                        "achieved": achieved, "male": male, "female": female,
                        "in_chart": "Y", "in_gender": "Y" if (male or female) else "N",
                    })
    return records


TITLE_DEFAULT = "Activity Target vs Achievement"

def slug(text):
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")[:20]


def seed_from_excel(records, pillars=PILLARS):
    """Auto first draft: each pillar's indicators split into two charts."""
    out = []
    for key, title in pillars.items():
        rows = [r for r in records if r["pillar"] == key]
        if not rows:
            continue
        n_charts = 2 if len(rows) > 1 else 1
        size = math.ceil(len(rows) / n_charts)
        charts = [rows[i:i + size] for i in range(0, len(rows), size)]
        gender = [r for r in rows if r["male"] or r["female"]]
        out.append({
            "key": key, "title": title,
            "charts": [{"title": TITLE_DEFAULT, "type": "column",
                        "rows": [(r["indicator"], r["target"], r["achieved"]) for r in ch]}
                       for ch in charts],
            "male": sum(r["male"] for r in gender),
            "female": sum(r["female"] for r in gender),
        })
    return out


def _walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(sh.shapes)
        else:
            yield sh


def seed_from_deck(deck_path, pillars=PILLARS):
    """Copy the charts of a finished deck: titles, types, indicators, numbers."""
    prs = Presentation(deck_path)
    out = []
    for slide in prs.slides:
        title = " ".join(slide.shapes.title.text_frame.text.replace("\v", " ")
                         .replace("\u200b", "").split()) \
            if slide.shapes.title is not None else ""
        m = re.search(r"PILLAR:\s*(.+)$", title, flags=re.I)
        if not m:
            continue
        ptitle = m.group(1).strip()
        key = next((k for k, v in pillars.items() if v.lower() == ptitle.lower()), slug(ptitle))
        entry = {"key": key, "title": ptitle, "charts": [], "male": 0, "female": 0}
        shapes = [sh for sh in _walk(slide.shapes) if getattr(sh, "has_chart", False) and sh.has_chart]
        for sh in sorted(shapes, key=lambda x: x.left):
            ch = sh.chart
            cats = [str(c) for c in ch.plots[0].categories]
            if ch.chart_type == XL_CHART_TYPE.PIE:
                vals = list(ch.plots[0].series[0].values) if len(ch.plots[0].series) else []
                if len(vals) >= 2 and [c.lower() for c in cats[:2]] == ["male", "female"]:
                    entry["male"], entry["female"] = vals[0] or 0, vals[1] or 0
                continue
            series = {s.name.lower(): list(s.values) for s in ch.plots[0].series}
            if "target" not in series or "achieved" not in series or not cats:
                continue
            ttl = " ".join(ch.chart_title.text_frame.text.split()) if ch.has_title else TITLE_DEFAULT
            kind = "bar" if ch.chart_type == XL_CHART_TYPE.BAR_CLUSTERED else "column"
            entry["charts"].append({
                "title": ttl or TITLE_DEFAULT, "type": kind,
                "rows": [(c.strip(), series["target"][i] or 0, series["achieved"][i] or 0)
                         for i, c in enumerate(cats)]})
        out.append(entry)
    return out


COLS = ["Pillar", "Pillar title", "Chart #", "Chart title", "Chart type (column/bar)",
        "Indicator", "Annual target", "Achieved", "Include (Y/N)"]


def save_review(data, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "chart_data"
    ws.append(COLS)
    for p in data:
        for n, ch in enumerate(p["charts"], start=1):
            for label, tgt, ach in ch["rows"]:
                ws.append([p["key"], p["title"], n, ch["title"], ch["type"],
                           label, tgt, ach, "Y"])
    gs = wb.create_sheet("gender")
    gs.append(["Pillar", "Male", "Female"])
    for p in data:
        gs.append([p["key"], p["male"], p["female"]])
    for sheet in (ws, gs):
        for c in sheet[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="9B2D1F")
        sheet.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHI", (12, 36, 8, 34, 14, 60, 13, 11, 12)):
        ws.column_dimensions[col].width = w
    wb.save(path)
    print(f"Chart data saved: {path}")


def load_review(path):
    wb = load_workbook(path, data_only=True)
    pillars, order = {}, []
    for row in wb["chart_data"].iter_rows(min_row=2, values_only=True):
        if not row or not row[5] or str(row[8] or "Y").strip().upper() != "Y":
            continue
        key, ptitle, num, ctitle, ctype, label, tgt, ach = row[:8]
        if key not in pillars:
            pillars[key] = {"key": key, "title": ptitle, "charts": {}, "male": 0, "female": 0}
            order.append(key)
        chart = pillars[key]["charts"].setdefault(
            int(num or 1), {"title": ctitle or TITLE_DEFAULT,
                            "type": str(ctype or "column").strip().lower(), "rows": []})
        chart["rows"].append((str(label), float(tgt or 0), float(ach or 0)))
    for row in wb["gender"].iter_rows(min_row=2, values_only=True):
        if row and row[0] in pillars:
            pillars[row[0]]["male"], pillars[row[0]]["female"] = float(row[1] or 0), float(row[2] or 0)
    result = []
    for k in order:
        p = pillars[k]
        p["charts"] = [p["charts"][n] for n in sorted(p["charts"])]
        result.append(p)
    return result


# =========================================================================== #
# 3. BUILDING THE DECK
# =========================================================================== #
def _style_font(font, size, bold=False, color=None):
    font.size = Pt(size)
    font.bold = bold
    font.name = FONT
    if color is not None:
        font.color.rgb = color


def _set_title(chart, text, size=CHART_FONT_SIZE):
    chart.has_title = True
    tf = chart.chart_title.text_frame
    tf.text = text
    _style_font(tf.paragraphs[0].runs[0].font, size, True, C_GREY)
    chart.chart_title.include_in_layout = False


def add_bar_chart(slide, spec, x, y, w, h):
    horizontal = spec["type"] == "bar"
    rows = spec["rows"]
    labels = [r[0] if horizontal else shorten(r[0]) for r in rows]
    data = CategoryChartData()
    data.categories = labels
    data.add_series("target", [r[1] for r in rows])
    data.add_series("achieved", [r[2] for r in rows])
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    chart = slide.shapes.add_chart(kind, x, y, w, h, data).chart
    _set_title(chart, spec["title"])

    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    _style_font(chart.legend.font, CHART_FONT_SIZE, True)

    plot = chart.plots[0]
    plot.gap_width = 60
    plot.overlap = -5
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = "#,##0", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    _style_font(dl.font, CHART_FONT_SIZE, True)
    for series, colour in zip(plot.series, (C_TARGET, C_ACHIEVED)):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = colour
        sdl = series.data_labels                      # series-level labels (keeps every bar labelled)
        sdl.show_value = True
        sdl.number_format, sdl.number_format_is_linked = "#,##0", False
        sdl.position = XL_LABEL_POSITION.OUTSIDE_END
        _style_font(sdl.font, CHART_FONT_SIZE, True)
        # '%' indicators show their value with a % sign
        for i, row in enumerate(rows):
            if row[0].lstrip().startswith("%"):
                value = row[1] if series.name == "target" else row[2]
                tf = series.points[i].data_label.text_frame
                tf.text = f"{value:g}%"
                _style_font(tf.paragraphs[0].runs[0].font, CHART_FONT_SIZE, True)
                series.points[i].data_label.position = XL_LABEL_POSITION.OUTSIDE_END

    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False
    cat = chart.category_axis
    _style_font(cat.tick_labels.font, CHART_FONT_SIZE, True, C_GREY)
    cat.format.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
    if horizontal:
        cat.reverse_order = True                      # first indicator at the top
    else:
        cat._element.get_or_add_txPr().bodyPr.set("rot", "-2700000")   # slanted labels
    return chart


def add_gender_pie(slide, male, female, x, y, w, h):
    data = ChartData()
    data.categories = ["Male", "Female"]
    data.add_series("Gender", [male, female])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.PIE, x, y, w, h, data).chart
    _set_title(chart, "Gender Distribution")
    chart.has_legend = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.show_percentage = True
    dl.show_category_name = True
    dl.show_value = False
    dl.number_format, dl.number_format_is_linked = "0%", False
    _style_font(dl.font, CHART_FONT_SIZE, True, RGBColor(0xFF, 0xFF, 0xFF))
    for i, colour in enumerate((C_TARGET, C_ACHIEVED)):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = colour
    return chart


def open_template(path):
    """Use the supplied deck as a theme/layout source, with its slides removed."""
    if os.path.exists(path):
        prs = Presentation(path)
        ids = prs.slides._sldIdLst
        for sld in list(ids):
            prs.part.drop_rel(sld.rId)
            ids.remove(sld)
        return prs
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def new_slide(prs, title_lines):
    layout = next((l for l in prs.slide_layouts if l.name == "Title and Content"),
                  prs.slide_layouts[min(5, len(prs.slide_layouts) - 1)])
    slide = prs.slides.add_slide(layout)
    for ph in list(slide.placeholders):          # drop empty body placeholders
        if ph.placeholder_format.type != PP_PLACEHOLDER.TITLE:
            ph._element.getparent().remove(ph._element)
    for ph in layout.placeholders:               # bring back the slide-number badge
        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            slide.shapes._spTree.append(copy.deepcopy(ph._element))

    title = slide.shapes.title
    title.left, title.top = Inches(0.8), Inches(0.1)
    title.width, title.height = prs.slide_width - Inches(1.6), Inches(1.15)
    tf = title.text_frame
    tf.text = "\v".join(title_lines)             # \v = line break inside a paragraph
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    for run in tf.paragraphs[0].runs:
        _style_font(run.font, 19, True)
    return slide


def build_presentation(data, output_path, template_path=TEMPLATE_PPTX):
    """One slide per pillar: up to two bar charts + gender pie."""
    prs = open_template(template_path)
    W = prs.slide_width
    top, height = Inches(1.35), Inches(5.3)

    for n, p in enumerate(data):
        charts = p["charts"][:2]
        if not charts:
            print(f"  No chart data for '{p['key']}' - skipped")
            continue
        lines = ([DECK_TITLE] if n == 0 else []) + [f"PILLAR: {p['title']}"]
        slide = new_slide(prs, lines)
        has_pie = (p["male"] + p["female"]) > 0

        if has_pie and len(charts) == 2:             # chart | pie | chart
            pie_w = Inches(3.3)
            side = int((W - pie_w - Inches(0.5)) / 2)
            add_bar_chart(slide, charts[0], Inches(0.25), top, side, height)
            add_gender_pie(slide, p["male"], p["female"], int(Inches(0.25) + side),
                           top + Inches(0.6), int(pie_w), Inches(3.4))
            add_bar_chart(slide, charts[1], int(W - Inches(0.25) - side), top, side, height)
        elif has_pie:                                # chart | pie
            add_bar_chart(slide, charts[0], Inches(0.4), top, int(W * 0.62), height)
            add_gender_pie(slide, p["male"], p["female"], int(W * 0.62 + Inches(0.6)),
                           top + Inches(0.6), int(W * 0.30), Inches(3.6))
        elif len(charts) == 2:                       # chart | chart
            half = int((W - Inches(0.8)) / 2)
            add_bar_chart(slide, charts[0], Inches(0.4), top, half, height)
            add_bar_chart(slide, charts[1], Inches(0.4) + half, top, half, height)
        else:
            add_bar_chart(slide, charts[0], Inches(1.5), top, int(W - Inches(3.0)), height)

        print(f"  {p['title']}: {len(charts)} chart(s), "
              f"{'with' if has_pie else 'no'} gender pie (M={p['male']:,.0f} F={p['female']:,.0f})")

    prs.save(output_path)
    print(f"\nDeck saved: {output_path}")
    return output_path


def make_highlights(folder=None, refresh=False, reparse=False, from_deck=False):
    folder = Path(folder) if folder else Path(__file__).resolve().parent
    excel_path = folder / EXCEL_NAME
    review_path = folder / REVIEW_NAME
    template = folder / TEMPLATE_PPTX

    if refresh or not excel_path.exists():
        print("Extracting tables from the pptx files ...")
        pptx_tables_to_excel(folder, EXCEL_NAME)        # your first function

    if from_deck:
        print(f"Seeding chart data from {template.name} ...")
        data = seed_from_deck(str(template))
        save_review(data, review_path)
    elif reparse or not review_path.exists():
        print(f"Parsing {excel_path.name} ...")
        data = seed_from_excel(extract_indicators(excel_path))
        save_review(data, review_path)
    else:
        print(f"Using your chart data: {review_path.name}  (use --reparse to rebuild it)")

    data = load_review(review_path)
    return build_presentation(data, str(folder / OUTPUT_PPTX), str(template))


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    make_highlights(args[0] if args else None,
                    refresh="--refresh" in sys.argv,
                    reparse="--reparse" in sys.argv,
                    from_deck="--from-deck" in sys.argv)
